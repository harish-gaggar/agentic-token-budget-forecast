#!/usr/bin/env python3
"""Build IEEE TAI submission files (double-anonymous compliant).

Outputs under submission/ieee-tai/output/:
  Anonymized_Main_Manuscript.docx
  Anonymized_Main_Manuscript_LaTeX.zip
  Main_Manuscript.pdf (preview)
  Title_Page.docx
  Conflict_of_Interest.docx
  Cover_Letter.docx
  Supplementary_Material_for_Review.docx
  SUBMISSION_CHECKLIST.md

Run from repo root:
  python submission/ieee-tai/build_submission.py
"""

from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
PKG = Path(__file__).resolve().parent
OUT = PKG / "output"
BUILD = PKG / ".build"
TEMPLATE = Path("/Users/hgaggar/Downloads/TAI_Word_Template.doc")
TEMPLATE_DOC = PKG / "IEEE_TAI_Word_Template.doc"
TEMPLATE_DOCX = PKG / "IEEE_TAI_Word_Template.docx"
PAPER_TXT = REPO / "paper.txt"
PAPER_IEEE_TEX = REPO / "paper_ieee.tex"
PAPER_IEEE_PDF = REPO / "paper_ieee.pdf"
PANDOC_IMAGE = "pandoc/extra:3.6"
PUBLIC_REPO_URL = "https://github.com/harish-gaggar/agentic-token-budget-forecast"

IEEE_FIGURES = [
    "results/fig_w20r.pdf",
    "results/fig_calibration.pdf",
    "results/fig_error_dist.pdf",
    "results/fig_per_type_w20r.pdf",
    "results/fig_production_savings.pdf",
    "results/fig_generic_agents.pdf",
]

TITLE = "Task-Aware Token Budget Forecasting for Agentic Workflows"

ABSTRACT = (
    "Agentic large language model workflows routinely consume far more tokens than "
    "single-turn chat because they interleave model calls, tools, and growing "
    "conversation state. Teams still size context reactively, which raises cost and "
    "risks context-window failures. This paper asks whether the task description "
    "available at dispatch time is sufficient to forecast per-stage token use before "
    "execution. The proposed Task-Aware Budget Forecaster (TABF) extracts lexical "
    "signals, classifies task type and complexity, and predicts a per-stage budget "
    "without any large language model call. A gradient-boosted regressor trained on "
    "simulated traces raises the within-twenty-percent rate from 65.4% to 74.1% on a "
    "six-hundred-task reproducible benchmark (McNemar p = 0.009). A complementary "
    "in-loop framework, ContextOptimizer, trims live message lists before each model "
    "call while preserving tool-call and response structure. Evaluations on public "
    "multi-turn trace fixtures and synthetic agent personas report up to 72% "
    "input-token reduction on held-out traces. The study contributes open benchmarks, "
    "evaluation harnesses, and reproducible measurements aimed at predictable "
    "agentic inference cost."
)

IMPACT = (
    "Enterprises are deploying agentic AI systems whose inference bills scale with "
    "conversation length and tool output, yet budgets are still set with static "
    "defaults. Accurate pre-flight token forecasts enable cost gates, model routing, "
    "and capacity planning before expensive multi-step runs start. In-loop trimming "
    "that preserves tool protocols can cut recurring input spend without adding "
    "another model call in the optimiser path. Together, forecasting plus structured "
    "compression address a growing share of AI operating expenditure and reduce "
    "out-of-context failures that interrupt customer-facing workflows. The artifacts "
    "support reproducible benchmarking so practitioners can compare methods on equal "
    "footing. Wider adoption could standardise how platforms report and govern token "
    "use across agent frameworks."
)

KEYWORDS = (
    "Agent-based systems; Natural language processing; Machine learning; "
    "Knowledge-based systems; Applications"
)


def _word_count(text: str) -> int:
    return len(text.split())


def _anonymize_body(text: str) -> str:
    """Light pass for double-anonymous review (TAI policy)."""
    text = re.sub(r"\bWe ask\b", "This work asks", text)
    text = re.sub(r"\bWe answer\b", "The proposed pipeline answers", text)
    text = re.sub(r"\bWe treat\b", "This work treats", text)
    text = re.sub(r"\bWe evaluate\b", "The evaluation", text)
    text = re.sub(r"\bWe are not aware\b", "Prior work has not", text)
    text = re.sub(r"\bOur starting point\b", "The starting point", text)
    text = re.sub(r"\bour\b", "the", text)
    text = re.sub(r"\bOur\b", "The", text)
    text = re.sub(r"\bwe\b", "the authors", text)
    text = re.sub(r"\bWe\b", "The paper", text)
    # Remove de-identifying links and release callouts in blind manuscript.
    text = re.sub(
        r"All code, CSVs, and seeds are released\.?\s*",
        "Reproducibility artifacts are described in the supplementary material.\n\n",
        text,
    )
    text = re.sub(
        r"released with this paper\.?\s*",
        "described in the supplementary material.\n\n",
        text,
    )
    text = re.sub(r"https?://\S+", "[URL removed for review]", text)
    text = re.sub(r"github\.com/\S+", "[repository URL withheld for review]", text)
    text = re.sub(r"Intuit Inc\.?", "[Affiliation withheld]", text)
    text = re.sub(r"Harish Gaggar", "Anonymous", text)
    text = re.sub(r"prod-ck-analytics-data-\d+", "[dataset withheld]", text)
    return text


def _load_body() -> str:
    raw = PAPER_TXT.read_text(encoding="utf-8")
    # Drop title line and abstract block from paper.txt export.
    parts = raw.split("\n\n", 2)
    if len(parts) >= 3 and parts[1].strip().lower() == "abstract":
        body = parts[2]
    else:
        body = raw
    # Remove LaTeX comment artifacts.
    body = re.sub(r"^%.*$", "", body, flags=re.MULTILINE)
    return _anonymize_body(body.strip())


def _set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(10)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        if name in doc.styles:
            h = doc.styles[name]
            h.font.name = "Times New Roman"
            h.font.color.rgb = RGBColor(0, 0, 0)


def _ensure_template_docx() -> Path:
    """IEEE TAI ships a .doc template; convert once to .docx for pandoc reference-doc."""
    if TEMPLATE_DOCX.exists():
        return TEMPLATE_DOCX
    source = TEMPLATE_DOC if TEMPLATE_DOC.exists() else TEMPLATE
    if not source.exists():
        raise SystemExit(
            f"Missing IEEE Word template. Copy TAI_Word_Template.doc to {TEMPLATE_DOC}"
        )
    subprocess.run(
        [
            "textutil",
            "-convert",
            "docx",
            "-output",
            str(TEMPLATE_DOCX),
            str(source),
        ],
        check=True,
    )
    return TEMPLATE_DOCX


def _apply_ieee_word_layout(doc: Document) -> None:
    """Approximate IEEE two-column letter layout (PDF remains the authoritative version)."""
    _set_normal_style(doc)

    for section in doc.sections:
        sect_pr = section._sectPr
        for el in sect_pr.findall(qn("w:cols")):
            sect_pr.remove(el)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "720")  # 0.5 in between columns
        sect_pr.append(cols)

        pg_mar = sect_pr.find(qn("w:pgMar"))
        if pg_mar is None:
            pg_mar = OxmlElement("w:pgMar")
            sect_pr.append(pg_mar)
        # 0.75 in margins (1080 twips)
        for side in ("top", "bottom", "left", "right"):
            pg_mar.set(qn(f"w:{side}"), "1080")

    for para in doc.paragraphs:
        if para.style and para.style.name and para.style.name.startswith("Heading"):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.name = "Times New Roman"
                run.bold = True
                run.italic = False


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text.strip("# ").strip(), level=level)


def _add_paragraphs_from_text(doc: Document, text: str) -> None:
    for block in re.split(r"\n\n+", text):
        block = block.strip()
        if not block:
            continue
        if block.startswith("## "):
            _add_heading(doc, block[2:], level=1)
        elif block.startswith("### "):
            _add_heading(doc, block[4:], level=2)
        elif block.startswith("[Figure") or block.startswith("[Table"):
            p = doc.add_paragraph(block)
            p.italic = True
        else:
            # Merge hard-wrapped lines within a paragraph.
            lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
            doc.add_paragraph(" ".join(lines))


# IEEE section numbers/titles from paper_ieee.aux (match PDF).
IEEE_SECTION_HEADINGS: dict[str, tuple[str, str]] = {
    "impact statement": ("I", "Impact Statement"),
    "introduction": ("II", "Introduction"),
    "related work": ("III", "Related Work"),
    "problem formulation": ("IV", "Problem Formulation"),
    "the tabf framework": ("V", "The TABF Framework"),
    "experimental setup": ("VI", "Experimental Setup"),
    "results": ("VII", "Results"),
    "contextoptimizer": ("VIII", "ContextOptimizer"),
    "trace and persona evaluation": ("IX", "Trace and Persona Evaluation"),
    "discussion": ("X", "Discussion"),
    "conclusion": ("XI", "Conclusion"),
}

IEEE_SUBSECTION_HEADINGS: dict[str, tuple[str, str]] = {
    "signal extractor and classifier": ("A", "Signal Extractor and Classifier"),
    "heuristic and gbt budget models": ("B", "Heuristic and GBT Budget Models"),
}

# Anchor phrase in body text -> table should appear immediately after this paragraph.
TABLE_PLACEMENT_ANCHORS: list[tuple[str, str]] = [
    ("tab:base", "Table I lists"),
    ("tab:main", "Table II reports"),
    ("tab:ablation", "Selected heuristic ablations"),
    ("tab:oracle", "(Table IV)"),
    ("tab:real", "(Table V)"),
    ("tab:prodcost", "Trace evaluation savings"),
    ("tab:personas", "Persona means"),
]

FIGURE_CAPTION_PREFIXES: list[tuple[str, str]] = [
    ("Within-20% rate by method", "Fig. 1"),
    ("Predicted vs.", "Fig. 2"),
    ("TABF+GBT W20R by latent task type", "Fig. 3"),
    ("Per-call input tokens before/after", "Fig. 4"),
    ("Persona benchmark: tokens before/after", "Fig. 5"),
]


def _prepare_tex_for_word(tex: str) -> str:
    """LaTeX tweaks so pandoc Word output matches paper_ieee.pdf body text."""
    # Custom \\paragraph spacing becomes literal '1.5ex -1em' in Word; drop it.
    tex = re.sub(
        r"% Run-in paragraph headings.*?\\makeatother\n",
        lambda _m: "\n\\makeatother\n",
        tex,
        count=1,
        flags=re.DOTALL,
    )
    return tex


def _write_word_source_tex(*, anonymized: bool = False) -> Path:
    """Same scientific content as paper_ieee.tex; optional blind author line."""
    tex = _prepare_tex_for_word(PAPER_IEEE_TEX.read_text(encoding="utf-8"))
    if anonymized:
        tex = re.sub(
            r"\\author\{[^}]*\}",
            r"\\author{Anonymous Author(s)}",
            tex,
            count=1,
        )
    BUILD.mkdir(parents=True, exist_ok=True)
    name = "main_anon_word.tex" if anonymized else "main_word.tex"
    path = BUILD / name
    path.write_text(tex, encoding="utf-8")
    return path


def _write_anon_tex(*, portal_standard: bool = False) -> Path:
    """Double-anonymous main manuscript source (no author affiliation)."""
    text = PAPER_IEEE_TEX.read_text(encoding="utf-8")
    text = re.sub(
        r"\\author\{[^}]*\}",
        r"\\author{Anonymous Author(s)}",
        text,
        count=1,
    )
    if portal_standard:
        # IEEE portal compiler expects stock IEEEtran abstract/keywords markup.
        text = re.sub(
            r"% Plain abstract/keywords labels.*?\\makeatother\n",
            "",
            text,
            count=1,
            flags=re.DOTALL,
        )
    BUILD.mkdir(parents=True, exist_ok=True)
    name = "main_anon_portal.tex" if portal_standard else "main_anon.tex"
    anon = BUILD / name
    anon.write_text(text, encoding="utf-8")
    return anon


def _pandoc_tex_to_docx(tex_path: Path, docx_path: Path, *, ieee_portal_layout: bool = False) -> None:
    ref_doc = _ensure_template_docx()
    rel_tex = tex_path.relative_to(REPO).as_posix()
    rel_out = docx_path.relative_to(REPO).as_posix()
    rel_ref = ref_doc.relative_to(REPO).as_posix()
    cmd = [
        "docker",
        "run",
        "--rm",
        "--platform",
        "linux/amd64",
        "-v",
        f"{REPO}:/data",
        "-w",
        "/data",
        PANDOC_IMAGE,
        rel_tex,
        "-o",
        rel_out,
        "--from=latex",
        "--to=docx",
        f"--reference-doc={rel_ref}",
        f"--resource-path=.:results:submission/ieee-tai/.build",
        "--bibliography=paper.bib",
        "--citeproc",
    ]
    subprocess.run(cmd, check=True)
    if ieee_portal_layout:
        return
    doc = Document(docx_path)
    _apply_ieee_word_layout(doc)
    doc.save(docx_path)


def _set_paragraph_text(paragraph, text: str) -> None:
    """Replace full paragraph text (pandoc may leave stray field/hyperlink nodes)."""
    element = paragraph._element
    for child in list(element):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            element.remove(child)
    paragraph.add_run(text)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def _parse_aux_refs(aux_path: Path) -> dict[str, str]:
    """Map label keys (tab:main, fig:w20r) to IEEE-style numbers from paper_ieee.aux."""
    refs: dict[str, str] = {}
    if not aux_path.exists():
        return refs
    for match in re.finditer(
        r"\\newlabel\{([^}]+)\}\{\{([^}]+)\}", aux_path.read_text(encoding="utf-8")
    ):
        refs[match.group(1)] = match.group(2)
    return refs


def _citeproc_table_arabic_to_roman() -> dict[str, str]:
    """Map pandoc citeproc 'Table 2'.. to IEEE Roman numerals from paper_ieee.tex refs."""
    refs = _parse_aux_refs(REPO / "paper_ieee.aux")
    order: list[str] = []
    for match in re.finditer(r"\\ref\{(tab:\w+)\}", PAPER_IEEE_TEX.read_text(encoding="utf-8")):
        key = match.group(1)
        if key not in order:
            order.append(key)
    # tab:base / tab:main become Table I / II via \\ref fix; citeproc numbers the rest 2..N.
    numeric_keys = [key for key in order if key not in ("tab:base", "tab:main")]
    mapping: dict[str, str] = {}
    for arabic, key in enumerate(numeric_keys, start=2):
        if key in refs:
            mapping[str(arabic)] = refs[key]
    return mapping


def _fix_citeproc_table_numbers(text: str) -> str:
    mapping = _citeproc_table_arabic_to_roman()
    for arabic in sorted(mapping, key=int, reverse=True):
        roman = mapping[arabic]
        text = text.replace(f"Table\u00a0{arabic}", f"Table {roman}")
        text = text.replace(f"Table {arabic}", f"Table {roman}")
    return text


def _fix_crossrefs(text: str, refs: dict[str, str]) -> str:
    for key, num in refs.items():
        bracket = f"[{key}]"
        if key.startswith("tab:"):
            label = f"Table {num}"
            text = re.sub(
                rf"Table\s*{re.escape(bracket)}",
                label,
                text,
                flags=re.IGNORECASE,
            )
            # Pandoc often leaves a duplicate bracket after the inline Table mention.
            text = text.replace(bracket, "")
        elif key.startswith("fig:"):
            label = f"Figure {num}"
            text = re.sub(
                rf"Figure\s*{re.escape(bracket)}",
                label,
                text,
                flags=re.IGNORECASE,
            )
            text = text.replace(bracket, "")
        elif key.startswith("sec:"):
            text = text.replace(bracket, f"Section {num}")
    return re.sub(r"  +", " ", text)


RUNIN_PARAGRAPH_LABELS = {
    "robustness.": "a",
    "forecasting.": "a",
    "optimisation.": "b",
    "composition.": "c",
}


def _merge_runin_paragraph_headings(doc: Document) -> None:
    """Pandoc splits \\paragraph{Title.} from its body; merge and prefix a), b), c)."""
    i = 0
    while i < len(doc.paragraphs) - 1:
        head = doc.paragraphs[i].text.strip().lower()
        if head in RUNIN_PARAGRAPH_LABELS:
            letter = RUNIN_PARAGRAPH_LABELS[head]
            title = doc.paragraphs[i].text.strip()
            body = doc.paragraphs[i + 1].text.strip()
            if body:
                _set_paragraph_text(
                    doc.paragraphs[i], f"{letter}) {title} {body}"
                )
                _remove_paragraph(doc.paragraphs[i + 1])
                continue
        i += 1


def _merge_label_with_body(doc: Document, label: str, prefix: str) -> None:
    """Merge a label paragraph with the following body (same text as LaTeX/PDF)."""
    target = label.lower().rstrip(":")
    for i in range(len(doc.paragraphs) - 1):
        head = doc.paragraphs[i].text.strip().lower().rstrip(":")
        if head == target:
            body = doc.paragraphs[i + 1].text
            _set_paragraph_text(doc.paragraphs[i], prefix + body)
            _remove_paragraph(doc.paragraphs[i + 1])
            return


def _restore_pdf_math_fragments(text: str) -> str:
    """Restore only empty math placeholders pandoc drops (wording from paper_ieee.tex)."""
    fixes = [
        ("reports held-out accuracy ().", "reports held-out accuracy (n = 185)."),
        (
            "reproducible benchmark (McNemar ). ContextOptimizer",
            "reproducible benchmark (McNemar p = 0.009). ContextOptimizer",
        ),
        ("65.4% heuristic () and more", "65.4% heuristic (p = 0.009) and more"),
        ("statistically tied (McNemar ) but lack", "statistically tied (McNemar p = 1.0) but lack"),
        ("McNemar tests (, exact", "McNemar tests (n = 185, exact"),
        (
            "Five-seed cross-validation yields W20R for TABF+GBT",
            "Five-seed cross-validation yields 71.7±4.0% W20R for TABF+GBT",
        ),
        ("Five-seed CV () and stress", "Five-seed CV (71.7±4.0%) and stress"),
        ("scales them by complexity (-).", "scales them by complexity (0.65–2.4×)."),
        ("scale columns by -).", "scale columns by 0.65–2.4×)."),
        ("scale factors or  local traces", "scale factors or ~100 local traces"),
        ("within 20% on a reproducible", "within ±20% on a reproducible"),
        (
            "The pipeline is and makes no LLM call",
            "The pipeline is O(|τ|) and makes no LLM call",
        ),
        (
            "The extractor maps to using regular expressions only",
            "The extractor maps τ to σ(τ) ∈ R⁸ using regular expressions only",
        ),
        (
            "rules on to assign and , disambiguating",
            "rules on σ to assign κ and δ, disambiguating",
        ),
        ("$O(|\tau|)TABF", "TABF"),
        ("$O(|\\tau|)TABF", "TABF"),
        ("τσ(τ) ∈ R⁸σκδThe", "The"),
        ("τσ(τ) ∈ R⁸σκδ The", "The"),
        ("0.652.4 × 0.652.4 ×The", "The"),
        ("0.652.4 × 0.652.4 × The", "The"),
    ]
    for old, new in fixes:
        text = text.replace(old, new)
    return text


def _fix_pandoc_garbage(text: str) -> str:
    """Remove LaTeX paragraph spacing and math debris pandoc leaves in body text."""
    text = re.sub(r"1\.5ex\s+-1em\s+", "", text)
    text = re.sub(
        r"^[\s~]*(?:[\d.]+\s*±\s*[\d.]+\s*%?\s*)?(?:x\s+)?(?=[A-Za-z])",
        "",
        text,
    )
    return _fix_stray_leading_symbols(text)


def _fix_stray_leading_symbols(text: str) -> str:
    """Drop orphan ±/+ pandoc leaves before words (e.g. Conclusion opening)."""
    text = re.sub(r"^[±+]\s*(?=[A-Za-z])", "", text)
    text = text.replace("±Agentic", "Agentic").replace("+Agentic", "Agentic")
    return text


def _add_author_abstract_spacing(doc: Document) -> None:
    """Blank line between author block and abstract (IEEE title block layout)."""
    if len(doc.paragraphs) > 1:
        doc.paragraphs[1].paragraph_format.space_after = Pt(12)
    for para in doc.paragraphs:
        if para.text.startswith("Abstract:"):
            para.paragraph_format.space_before = Pt(6)
            break


def _relocate_table_after_anchor(doc: Document, table_el, anchor: str) -> None:
    """Move a Word table so it immediately follows the paragraph citing it."""
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    anchor_el = None
    for child in list(body):
        if child.tag != qn("w:p"):
            continue
        if anchor in Paragraph(child, doc).text:
            anchor_el = child
    if anchor_el is None:
        return
    parent = table_el.getparent()
    if parent is not None:
        parent.remove(table_el)
    anchor_el.addnext(table_el)


def _place_tables_near_citations(doc: Document) -> None:
    """Reorder tables next to the sentences that cite them (reviewer clarity)."""
    tables = [el for el in list(doc.element.body) if el.tag == qn("w:tbl")]
    for (_key, anchor), table_el in zip(TABLE_PLACEMENT_ANCHORS, tables):
        _relocate_table_after_anchor(doc, table_el, anchor)


def _remove_index_terms(doc: Document) -> None:
    """Portal abstract field covers keywords; drop Index Terms block from Word body."""
    for para in list(doc.paragraphs):
        if para.text.strip().lower().startswith("index terms"):
            _remove_paragraph(para)


def _remove_duplicate_paragraphs(doc: Document) -> None:
    """Remove back-to-back identical paragraphs (e.g. duplicated table titles)."""
    index = 0
    while index < len(doc.paragraphs) - 1:
        left = doc.paragraphs[index]
        right = doc.paragraphs[index + 1]
        left_text = left.text.strip()
        right_text = right.text.strip()
        if not left_text or left_text != right_text:
            if left_text.lower() != right_text.lower():
                index += 1
                continue
        left_style = left.style.name if left.style else ""
        right_style = right.style.name if right.style else ""
        if left_style in ("Table Caption", "Image Caption"):
            _remove_paragraph(right)
        elif right_style in ("Table Caption", "Image Caption"):
            _remove_paragraph(left)
        else:
            _remove_paragraph(right)
        # re-check same index in case of triple duplicate


def _label_tables_and_figures(doc: Document, refs: dict[str, str]) -> None:
    """Prefix pandoc caption paragraphs once (no extra inserted titles)."""
    table_caption_index = 0
    for para in doc.paragraphs:
        if not para.style or para.style.name != "Table Caption":
            continue
        if table_caption_index >= len(TABLE_PLACEMENT_ANCHORS):
            break
        key, _anchor = TABLE_PLACEMENT_ANCHORS[table_caption_index]
        roman = refs.get(key, "")
        text = re.sub(r"^TABLE\s+([IVX]+|\d+)\.\s*", "", para.text.strip(), flags=re.I)
        text = re.sub(r"^Table\s+([IVX]+|\d+)\.\s*", "", text, flags=re.I)
        if roman:
            _set_paragraph_text(para, f"Table {roman}. {text}")
        table_caption_index += 1

    for para in doc.paragraphs:
        if not para.style or para.style.name != "Image Caption":
            continue
        text = para.text.strip()
        if text.lower().startswith("fig."):
            continue
        for prefix, fig_label in FIGURE_CAPTION_PREFIXES:
            if text.startswith(prefix):
                _set_paragraph_text(para, f"{fig_label}. {text}")
                break


def _fix_conclusion_leading_symbol(doc: Document) -> None:
    """Remove orphan ±/+ before the Conclusion paragraph (separate run or paragraph)."""
    for para in list(doc.paragraphs):
        if re.fullmatch(r"[±+−\-=\s]+", para.text.strip()):
            _remove_paragraph(para)

    for index, para in enumerate(doc.paragraphs):
        if para.text.strip() != "XI. Conclusion":
            continue
        next_index = index + 1
        while next_index < len(doc.paragraphs):
            nxt = doc.paragraphs[next_index]
            if nxt.style and nxt.style.name == "Bibliography":
                break
            stripped = nxt.text.strip()
            if stripped in ("±", "+", "−", "-") or re.fullmatch(r"[±+−\-]+", stripped):
                _remove_paragraph(nxt)
                continue
            if "Agentic token budgeting" in stripped or stripped.startswith("Agentic"):
                cleaned = re.sub(r"^[\s±+−\-=]+\s*", "", nxt.text)
                cleaned = cleaned.replace("±Agentic", "Agentic").replace("+Agentic", "Agentic")
                if cleaned != nxt.text:
                    _set_paragraph_text(nxt, cleaned)
                for run in list(nxt.runs):
                    if run.text.strip() in ("±", "+", "−", "-", ""):
                        run._element.getparent().remove(run._element)
                    elif run.text.startswith("±") or run.text.startswith("+"):
                        run.text = re.sub(r"^[±+]+\s*", "", run.text)
                break
            next_index += 1
        break


def _number_bibliography_entries(doc: Document) -> None:
    """IEEE-style numbered references [1], [2], ... for Word bibliography paragraphs."""
    index = 0
    for para in doc.paragraphs:
        if not para.style or para.style.name != "Bibliography":
            continue
        index += 1
        text = para.text.strip()
        if re.match(r"^\[\d+\]\s", text):
            continue
        _set_paragraph_text(para, f"[{index}] {text}")
        para.paragraph_format.space_after = Pt(6)
        for run in para.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            _clear_italic(run)


def _set_bold_prefix(paragraph, prefix: str) -> None:
    """Bold label only (e.g. Abstract:), body regular — matches PDF."""
    if not paragraph.text.startswith(prefix):
        return
    body = paragraph.text[len(prefix) :]
    element = paragraph._element
    for child in list(element):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            element.remove(child)
    label_run = paragraph.add_run(prefix)
    label_run.bold = True
    label_run.font.name = "Times New Roman"
    label_run.font.size = Pt(10)
    if body:
        body_run = paragraph.add_run(body)
        body_run.bold = False
        body_run.font.name = "Times New Roman"
        body_run.font.size = Pt(10)


def _set_runin_heading_prefix(paragraph) -> None:
    """Format 'a) Forecasting.' run-in headings like the PDF."""
    match = re.match(r"^([a-z]\))\s+([A-Za-z][^.]*\.)\s+(.*)$", paragraph.text, re.DOTALL)
    if not match:
        return
    label = f"{match.group(1)} {match.group(2)} "
    body = match.group(3)
    element = paragraph._element
    for child in list(element):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            element.remove(child)
    head = paragraph.add_run(label)
    head.bold = True
    _clear_italic(head)
    head.font.name = "Times New Roman"
    head.font.size = Pt(10)
    if body:
        rest = paragraph.add_run(body)
        rest.bold = False
        _clear_italic(rest)
        rest.font.name = "Times New Roman"
        rest.font.size = Pt(10)
    paragraph.style = "Normal"


def _clear_italic(run) -> None:
    """Remove italic at API and OOXML level (pandoc Heading styles leave w:i)."""
    run.italic = False
    r_pr = run._element.find(qn("w:rPr"))
    if r_pr is None:
        return
    italic_el = r_pr.find(qn("w:i"))
    if italic_el is not None:
        r_pr.remove(italic_el)


def _demote_pandoc_heading_paragraphs(doc: Document) -> None:
    """\\paragraph run-ins exported as Heading 4 render italic in Word; use Normal."""
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 4":
            para.style = doc.styles["Normal"]


def _apply_section_heading_spacing(doc: Document) -> None:
    """One blank line before each Roman section heading (consistent vertical rhythm)."""
    known = {f"{num}. {title}" for num, title in IEEE_SECTION_HEADINGS.values()}
    for para in doc.paragraphs:
        text = para.text.strip()
        if text not in known:
            continue
        fmt = para.paragraph_format
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(0)


def _force_plain_body_typography(doc: Document) -> None:
    """No italic except table/figure captions (IEEE caption style)."""
    keep_italic_styles = {"Table Caption", "Image Caption"}
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        if style_name in keep_italic_styles:
            continue
        for run in para.runs:
            _clear_italic(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        _clear_italic(run)


def _flatten_hyperlinks(root) -> None:
    """Remove Word hyperlink wrappers (blue text) while keeping citation text."""
    for hyperlink in list(root.iter(qn("w:hyperlink"))):
        parent = hyperlink.getparent()
        if parent is None:
            continue
        index = list(parent).index(hyperlink)
        for child in list(hyperlink):
            if child.tag == qn("w:r"):
                parent.insert(index, child)
                index += 1
        parent.remove(hyperlink)


def _force_black_body_text(doc: Document) -> None:
    """All body text black; no hyperlink blue or citeproc highlight."""
    _flatten_hyperlinks(doc.element.body)
    black = RGBColor(0, 0, 0)
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.color.rgb = black
            run.font.highlight_color = None
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _flatten_hyperlinks(cell._element)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = black
                        run.font.highlight_color = None


def _finalize_word_typography(doc: Document) -> None:
    _demote_pandoc_heading_paragraphs(doc)
    _apply_section_heading_spacing(doc)
    _force_plain_body_typography(doc)
    _force_black_body_text(doc)


def _normalize_subsection_headings(doc: Document) -> None:
    """Subsections use A., B. (not V-A / V-B) for Word readability."""
    for para in doc.paragraphs:
        raw = para.text.strip()
        head = raw.lower()
        if head in IEEE_SUBSECTION_HEADINGS:
            num, title = IEEE_SUBSECTION_HEADINGS[head]
        elif head.startswith("v-a.") or head.startswith("v-b."):
            tail = raw.split(".", 1)[1].strip().lower()
            if tail not in IEEE_SUBSECTION_HEADINGS:
                continue
            num, title = IEEE_SUBSECTION_HEADINGS[tail]
        else:
            continue
        para.style = doc.styles["Normal"]
        _set_paragraph_text(para, f"{num}. {title}")
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            run.bold = True
            _clear_italic(run)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(0)


def _normalize_ieee_section_heading(paragraph, doc: Document) -> None:
    """Pandoc Heading 1/2 -> left-aligned 10pt bold section labels (IEEE style)."""
    head = paragraph.text.strip().lower()
    if head in IEEE_SUBSECTION_HEADINGS:
        num, title = IEEE_SUBSECTION_HEADINGS[head]
    elif head in IEEE_SECTION_HEADINGS:
        num, title = IEEE_SECTION_HEADINGS[head]
    else:
        return
    paragraph.style = doc.styles["Normal"]
    _set_paragraph_text(paragraph, f"{num}. {title}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.bold = True
        _clear_italic(run)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)


def _format_word_from_pandoc(doc: Document) -> None:
    """In-place formatting only; wording matches paper_ieee.tex / PDF."""
    refs = _parse_aux_refs(REPO / "paper_ieee.aux")

    if doc.paragraphs:
        title_p = doc.paragraphs[0]
        if TITLE.split()[0] in title_p.text:
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in title_p.runs:
                run.bold = True
                run.font.size = Pt(14)
    if len(doc.paragraphs) > 1:
        doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _merge_label_with_body(doc, "abstract", "Abstract: ")
    _remove_index_terms(doc)
    _merge_runin_paragraph_headings(doc)

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        if style_name.startswith("Heading"):
            _normalize_ieee_section_heading(para, doc)

    for para in doc.paragraphs:
        raw = para.text
        if not raw.strip():
            continue
        cleaned = raw.replace("\u2014", "-").replace("\u2013", "-")
        cleaned = _fix_pandoc_garbage(cleaned)
        cleaned = _restore_pdf_math_fragments(cleaned)
        cleaned = _fix_crossrefs(cleaned, refs)
        cleaned = _fix_citeproc_table_numbers(cleaned)
        if cleaned != raw:
            _set_paragraph_text(para, cleaned)

    for para in doc.paragraphs:
        _set_runin_heading_prefix(para)

    for para in doc.paragraphs:
        raw = para.text
        if not raw.strip():
            continue
        cleaned = _restore_pdf_math_fragments(raw)
        cleaned = _fix_citeproc_table_numbers(cleaned)
        if cleaned != raw:
            _set_paragraph_text(para, cleaned)
            _set_runin_heading_prefix(para)

    for para in doc.paragraphs:
        _set_bold_prefix(para, "Abstract: ")

    _normalize_subsection_headings(doc)
    _add_author_abstract_spacing(doc)
    _place_tables_near_citations(doc)
    _label_tables_and_figures(doc, refs)
    _remove_duplicate_paragraphs(doc)
    _fix_conclusion_leading_symbol(doc)
    _remove_duplicate_paragraphs(doc)
    _number_bibliography_entries(doc)
    _finalize_word_typography(doc)


def build_main_manuscript_word_portal(
    tex_path: Path,
    *,
    output_name: str = "Main_Manuscript.docx",
) -> Path:
    """Word export from same LaTeX source as paper_ieee.pdf (figures stay linked)."""
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / output_name
    print(f"Building Word from {tex_path.name} -> {output_name}...")
    _pandoc_tex_to_docx(tex_path, out, ieee_portal_layout=True)

    doc = Document(out)
    _set_normal_style(doc)
    _format_word_from_pandoc(doc)
    doc.save(out)
    return out


def build_anonymized_latex_zip() -> Path:
    """Zip for portal option: Anonymized Main Document - LaTeX File."""
    anon = _write_anon_tex(portal_standard=True)
    zip_path = OUT / "Anonymized_Main_Manuscript_LaTeX.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(anon, "main.tex")
        zf.write(REPO / "paper.bib", "paper.bib")
        for rel in IEEE_FIGURES:
            path = REPO / rel
            if not path.exists():
                raise SystemExit(f"Missing figure for LaTeX zip: {path}")
            zf.write(path, rel)
    return zip_path


def sync_main_manuscript_pdf() -> Path:
    """Copy the LaTeX-built IEEE PDF into the submission bundle."""
    if not PAPER_IEEE_PDF.exists():
        raise SystemExit("Missing paper_ieee.pdf. Build it first:\n  make ieee-blind-pdf")
    source = PAPER_IEEE_PDF
    dest = OUT / "Main_Manuscript.pdf"
    shutil.copy2(source, dest)
    return dest


def build_main_manuscript() -> Path:
    """Anonymized Word + LaTeX zip for IEEE TAI portal (PDF is layout reference only)."""
    if not PAPER_IEEE_TEX.exists():
        raise SystemExit(f"Missing {PAPER_IEEE_TEX}")

    sync_main_manuscript_pdf()
    _ensure_template_docx()

    # Word matching paper_ieee.pdf (author + body from paper_ieee.tex).
    word_tex = _write_word_source_tex(anonymized=False)
    build_main_manuscript_word_portal(word_tex, output_name="Main_Manuscript.docx")

    # Blind Word for IEEE portal upload only.
    anon_word_tex = _write_word_source_tex(anonymized=True)
    build_main_manuscript_word_portal(
        anon_word_tex, output_name="Anonymized_Main_Manuscript.docx"
    )

    print("Packaging anonymized LaTeX zip (standard IEEEtran abstract markup)...")
    build_anonymized_latex_zip()
    return OUT / "Main_Manuscript.docx"


def build_title_page() -> Path:
    if TEMPLATE.exists():
        path = OUT / "Title_Page.docx"
        shutil.copy2(TEMPLATE, path)
        doc = Document(path)
        # Clear sample paragraphs after title block by rebuilding minimal title page.
        doc = Document()
    else:
        doc = Document()
    _set_normal_style(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(14)

    doc.add_paragraph()
    doc.add_paragraph("Harish Gaggar")
    doc.add_paragraph("Intuit Inc.")
    doc.add_paragraph("Mountain View, CA, USA")
    doc.add_paragraph("Email: harish_gaggar@intuit.com")

    doc.add_heading("Corresponding Author", level=2)
    doc.add_paragraph(
        "Harish Gaggar, harish_gaggar@intuit.com, Intuit Inc., Mountain View, CA, USA"
    )

    doc.add_heading("Acknowledgments", level=2)
    doc.add_paragraph(
        "The authors thank colleagues who reviewed early drafts of this work."
    )

    path = OUT / "Title_Page.docx"
    doc.save(path)
    return path


def build_coi() -> Path:
    doc = Document()
    _set_normal_style(doc)
    doc.add_heading("Conflict of Interest Disclosure", level=1)
    doc.add_paragraph(
        "I declare the following interest relevant to this submission."
    )
    doc.add_paragraph(
        "I am employed by Intuit Inc. Experiments in the manuscript use public "
        "benchmarks, synthetic personas, and released evaluation traces; no "
        "proprietary or employer-internal datasets are described in the paper. "
        "I have no other financial conflicts of interest."
    )
    path = OUT / "Conflict_of_Interest.docx"
    doc.save(path)
    return path


def build_cover_letter() -> Path:
    doc = Document()
    _set_normal_style(doc)
    doc.add_paragraph("Date: May 25, 2025")
    doc.add_paragraph()
    doc.add_paragraph("Dear Editor-in-Chief, IEEE Transactions on Artificial Intelligence,")
    doc.add_paragraph()
    doc.add_paragraph(
        f"Please consider the attached manuscript, \"{TITLE},\" for publication as an "
        "Original Research Regular Manuscript in IEEE TAI."
    )
    doc.add_paragraph(
        "The paper introduces Task-Aware Budget Forecasting (TABF), a lightweight "
        "pre-flight pipeline that predicts per-stage token consumption from a task "
        "description without calling a large language model, and ContextOptimizer, "
        "an in-loop trimming framework that preserves tool-call/response structure. "
        "Contributions include a seeded six-hundred-task benchmark with byte-reproducible "
        "evaluation, statistically tested forecasting gains, optimizer measurements on "
        "public trace fixtures and synthetic personas, and open artifacts for replication."
    )
    doc.add_paragraph(
        "This manuscript is original, is not under consideration elsewhere, and has "
        "not been published previously."
    )
    doc.add_paragraph()
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph("Harish Gaggar")
    doc.add_paragraph("Intuit Inc.")
    doc.add_paragraph("harish_gaggar@intuit.com")
    path = OUT / "Cover_Letter.docx"
    doc.save(path)
    return path


SUPPLEMENTARY_CODE_DIRS = ("tabf", "benchmark", "experiments")
SUPPLEMENTARY_ROOT_FILES = ("requirements.txt", "Makefile")
SUPPLEMENTARY_SCRIPT_FILES = ("scripts/verify.sh",)
SUPPLEMENTARY_RESULTS = (
    "results/manifest.json",
    "results/main_results.csv",
    "results/significance.csv",
    "results/stratified_results.csv",
    "results/ablation_results.csv",
    "results/classifier_accuracy.json",
    "results/generic_agents.csv",
    "results/generic_agents_layers.csv",
    "results/production_token_savings_summary.csv",
    "results/production_token_savings_samples.csv",
    "results/optimizer_microbench.csv",
    "results/corpus.json",
    "results/traces.json",
    *IEEE_FIGURES,
)


def _zip_path(zf: zipfile.ZipFile, abs_path: Path, arcname: str | None = None) -> None:
    arc = arcname or abs_path.relative_to(REPO).as_posix()
    zf.write(abs_path, arc)


def _zip_tree(zf: zipfile.ZipFile, directory: Path) -> None:
    for path in sorted(directory.rglob("*")):
        if path.is_dir():
            continue
        if path.name == "__pycache__" or path.suffix in {".pyc", ".pyo"}:
            continue
        _zip_path(zf, path)


def build_supplementary_latex_zip() -> Path:
    """Artifacts zip for portal 'LaTeX Supplementary' / Supplementary Material."""
    BUILD.mkdir(parents=True, exist_ok=True)
    supp_tex = BUILD / "supplementary.tex"
    supp_tex.write_text(
        rf"""\documentclass[10pt]{{article}}
\usepackage[margin=0.75in]{{geometry}}
\usepackage{{url}}
\begin{{document}}
\noindent\textbf{{Supplementary Material (reproducibility).}}

\noindent\textbf{{Public repository:}} \url{{{PUBLIC_REPO_URL}}}

\noindent Clone the repository and run \texttt{{./scripts/verify.sh}} from the
repository root (Python~3.10--3.12, \texttt{{requirements.txt}}). Expected
synthetic benchmark: TABF\_gbt W20R $\approx$ 74.05\% in \texttt{{results/manifest.json}}.

\noindent This archive mirrors the repository: data tables, figure PDFs, and code
to reproduce the synthetic benchmark and figures reported in the manuscript.
Install dependencies from \texttt{{requirements.txt}}, then run
\texttt{{scripts/verify.sh}} from the bundle root.
\end{{document}}
""",
        encoding="utf-8",
    )

    zip_path = OUT / "Anonymized_Supplementary_LaTeX.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(supp_tex, "supplementary.tex")
        for name in SUPPLEMENTARY_ROOT_FILES:
            _zip_path(zf, REPO / name)
        for name in SUPPLEMENTARY_SCRIPT_FILES:
            _zip_path(zf, REPO / name)
        for dirname in SUPPLEMENTARY_CODE_DIRS:
            _zip_tree(zf, REPO / dirname)
        for rel in SUPPLEMENTARY_RESULTS:
            path = REPO / rel
            if not path.exists():
                raise SystemExit(f"Missing supplementary artifact: {path}")
            _zip_path(zf, path)
    return zip_path


def build_supplementary() -> Path:
    doc = Document()
    _set_normal_style(doc)
    doc.add_heading("Supplementary Material for Review", level=1)
    doc.add_heading("Public repository", level=2)
    doc.add_paragraph(
        f"Source code, data, and reproduction scripts: {PUBLIC_REPO_URL}"
    )
    doc.add_paragraph(
        "Quick start: git clone the repository, then from the repository root run "
        "./scripts/verify.sh (Python 3.10–3.12; see requirements.txt). "
        "Expected: TABF_gbt W20R ≈ 74.05% in results/manifest.json."
    )
    doc.add_paragraph(
        "The accompanying archive Anonymized_Supplementary_LaTeX.zip contains the "
        "same artifacts for offline review."
    )
    doc.add_heading("Environment", level=2)
    doc.add_paragraph("Python 3.10–3.12; dependencies in requirements.txt.")
    doc.add_heading("Reproduce synthetic benchmark (TABF)", level=2)
    doc.add_paragraph("make venv")
    doc.add_paragraph("make experiments")
    doc.add_paragraph("Expected: TABF_gbt W20R ≈ 74.05% in results/manifest.json")
    doc.add_heading("Reproduce figures from checked-in CSVs", level=2)
    doc.add_paragraph("make production")
    doc.add_paragraph("python -m experiments.plot_generic_agents --out results")
    doc.add_heading("Reviewer-evidence scripts", level=2)
    doc.add_paragraph("make multiseed stagewise calibration stress tabf-replay")
    doc.add_paragraph("make stagewise-gbt real-trace quality-outcomes compression-baseline oracle-labels")
    doc.add_heading("Supplementary archive contents", level=2)
    doc.add_paragraph(
        "Upload Anonymized_Supplementary_LaTeX.zip with: supplementary.tex, "
        "requirements.txt, Makefile, scripts/verify.sh, tabf/, benchmark/, "
        "experiments/, results/ tables and figure PDFs listed in the manuscript."
    )
    path = OUT / "Supplementary_Material_for_Review.docx"
    doc.save(path)
    return path


def build_checklist() -> Path:
    text = f"""# IEEE TAI submission checklist

Generated for: **{TITLE}**

Portal: https://ieee.atyponrex.com/journal/tai-ieee

## Required uploads (from portal)

| File | Format | Output file |
|------|--------|-------------|
| Main Manuscript (PDF match, local) | Word | `Main_Manuscript.docx` (author + same text as `Main_Manuscript.pdf`) |
| Main Manuscript (anonymized, portal) | **LaTeX zip** (best) or Word | `Anonymized_Main_Manuscript_LaTeX.zip` or `Anonymized_Main_Manuscript.docx` |
| Title Page | .docx only | `output/Title_Page.docx` |
| Conflict of Interest | .docx / .pdf / .rtf | `output/Conflict_of_Interest.docx` |
| Supplementary (LaTeX zip) | .zip | `Anonymized_Supplementary_LaTeX.zip` |
| Supplementary (instructions) | .docx | `Supplementary_Material_for_Review.docx` |

## Double-anonymous (mandatory)

- Main manuscript has **no** author names, affiliations, emails, acknowledgments, or funding.
- **No GitHub URLs** in the main manuscript (stay anonymized). Repository URL is in supplementary materials only.
- Use third person for prior work; avoid identifiable self-citations.
- Title page holds all identifying information (not sent to reviewers).

## Before you upload — manual steps

1. **Page limit:** IEEE TAI Regular = **10 pages** (+5 extra at fee). NeurIPS draft ≈15 pages — **condense** or budget extra pages.
2. **Main file (portal):** Upload **`Anonymized_Main_Manuscript_LaTeX.zip`** (recommended) or **`Anonymized_Main_Manuscript.docx`** for blind review. Use **`Main_Manuscript.docx`** only to compare against `Main_Manuscript.pdf` locally (includes author; not for blind upload).
3. **Tables:** Re-create from CSVs or export from PDF; plain-text export omits layout.
4. **Abstract:** {len(ABSTRACT.split())} words (limit 250). Impact: {len(IMPACT.split())} words (target 100–150).
5. **Keywords:** Pick 3–6 from TAI dropdown closest to: agent-based systems, NLP, machine learning.
6. **ORCID:** Required for all authors in portal.
7. **AI disclosure:** Complete only if required in the IEEE TAI submission portal (not on the title page).
8. **arXiv:** If posted, declare in portal + cover letter; explain differences if extending prior work.
9. **COI:** Upload `Conflict_of_Interest.docx` if the portal asks for a file; in the portal COI questions, disclose employment (do not select “none to disclose” if that would hide employer affiliation).
10. **Supplementary:** Upload `Anonymized_Supplementary_LaTeX.zip` if the portal has a LaTeX supplementary slot; also upload `Supplementary_Material_for_Review.docx` when required.
11. **Optional:** Upload `Cover_Letter.docx` (editor only, not reviewers).

## Regenerate files

```bash
python submission/ieee-tai/build_submission.py
```

## Verify science locally

```bash
./scripts/verify.sh
```
"""
    path = OUT / "SUBMISSION_CHECKLIST.md"
    path.write_text(text, encoding="utf-8")
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    assert _word_count(ABSTRACT) <= 250, f"Abstract too long: {_word_count(ABSTRACT)}"
    assert 100 <= _word_count(IMPACT) <= 155, f"Impact length: {_word_count(IMPACT)}"

    files = [
        build_main_manuscript(),
        build_title_page(),
        build_coi(),
        build_cover_letter(),
        build_supplementary(),
        build_supplementary_latex_zip(),
        build_checklist(),
    ]
    print("Wrote:")
    for f in files:
        print(f"  {f}")
    print(f"\nAbstract words: {_word_count(ABSTRACT)}")
    print(f"Impact words: {_word_count(IMPACT)}")


if __name__ == "__main__":
    main()
